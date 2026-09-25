"""
ExcelLedger - keeps a billing_ledger.xlsx workbook with THREE sheets, similar
in purpose to how Apache POI is used to read/write .xlsx files in Java:

  - "Rooms"   : one row per room (number, type, base rent, max occupancy)
  - "Tenants" : one row per tenant (name, contact number, assigned room)
  - "Billing" : one row per generated bill, updated in place when paid

Because all three sheets live in the same file, the whole application state
(rooms, tenants, and billing history) can be reloaded from this one file
the next time the program starts.
"""

import os
from openpyxl import Workbook, load_workbook


class ExcelLedger:
    ROOM_HEADERS = ["Room Number", "Type", "Base Rent", "Max Occupancy"]
    TENANT_HEADERS = ["Name", "Contact Number", "Room Number"]
    BILLING_HEADERS = ["Tenant", "Room", "Month", "Rent Share", "Utility Share",
                        "Total Amount", "Due Date", "Amount Paid", "Date Paid",
                        "Balance", "Status"]

    def __init__(self, filepath="billing_ledger.xlsx"):
        self.filepath = filepath

        if os.path.exists(filepath):
            self.workbook = load_workbook(filepath)
        else:
            self.workbook = Workbook()
            self.workbook.remove(self.workbook.active)  # drop the default blank sheet

        self.rooms_sheet = self._get_or_create_sheet("Rooms", self.ROOM_HEADERS)
        self.tenants_sheet = self._get_or_create_sheet("Tenants", self.TENANT_HEADERS)
        self.billing_sheet = self._get_or_create_sheet("Billing", self.BILLING_HEADERS)
        self.workbook.save(self.filepath)

    def _get_or_create_sheet(self, name, headers):
        if name in self.workbook.sheetnames:
            return self.workbook[name]
        sheet = self.workbook.create_sheet(name)
        sheet.append(headers)
        return sheet

    # ------------------------------------------------------------------
    # Rooms
    # ------------------------------------------------------------------
    def add_room(self, room):
        self.rooms_sheet.append([
            room.room_number,
            type(room).__name__,
            room.base_rent,
            room.get_max_occupancy(),
        ])
        self.workbook.save(self.filepath)

    def load_rooms(self):
        rooms_data = []
        for row in self.rooms_sheet.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            rooms_data.append({
                "room_number": row[0],
                "type": row[1],
                "base_rent": row[2],
                "max_occupancy": row[3],
            })
        return rooms_data

    # ------------------------------------------------------------------
    # Tenants
    # ------------------------------------------------------------------
    def add_tenant(self, tenant, room):
        self.tenants_sheet.append([
            tenant.name,
            tenant.contact_number,
            room.room_number,
        ])
        self.workbook.save(self.filepath)

    def load_tenants(self):
        tenants_data = []
        for row in self.tenants_sheet.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            tenants_data.append({
                "name": row[0],
                "contact_number": row[1],
                "room_number": row[2],
            })
        return tenants_data

    # ------------------------------------------------------------------
    # Billing
    # ------------------------------------------------------------------
    def add_bill(self, bill):
        """Appends a new row for a bill. Called once per bill, each time
        bills are generated for a month."""
        room_number = bill.tenant.assigned_room.room_number if bill.tenant.assigned_room else ""
        self.billing_sheet.append([
            bill.tenant.name,
            room_number,
            bill.month,
            bill.rent_share,
            bill.utility_share,
            bill.get_total_amount(),
            bill.due_date,
            0.0,                        # Amount Paid starts at 0
            "",                         # Date Paid, filled in once paid
            bill.get_total_amount(),    # Balance starts equal to the total
            "PAID" if bill.paid else "UNPAID",
        ])
        self.workbook.save(self.filepath)

    def record_payment(self, bill, payment):
        """Finds this bill's row (by tenant + room + month) and subtracts
        the payment amount from its balance."""
        room_number = bill.tenant.assigned_room.room_number if bill.tenant.assigned_room else ""

        for row in self.billing_sheet.iter_rows(min_row=2):
            tenant_cell, room_cell, month_cell = row[0], row[1], row[2]
            if (tenant_cell.value == bill.tenant.name
                    and room_cell.value == room_number
                    and month_cell.value == bill.month):

                amount_paid_cell = row[7]
                date_paid_cell = row[8]
                balance_cell = row[9]
                status_cell = row[10]

                new_amount_paid = (amount_paid_cell.value or 0) + payment.amount_paid
                new_balance = bill.get_total_amount() - new_amount_paid

                amount_paid_cell.value = new_amount_paid
                date_paid_cell.value = payment.date_paid
                balance_cell.value = new_balance
                status_cell.value = "PAID" if bill.paid else "UNPAID"
                break

        self.workbook.save(self.filepath)

    def load_bills(self):
        bills_data = []
        for row in self.billing_sheet.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            bills_data.append({
                "tenant_name": row[0],
                "room_number": row[1],
                "month": row[2],
                "rent_share": row[3],
                "utility_share": row[4],
                "total_amount": row[5],
                "due_date": row[6],
                "amount_paid": row[7],
                "date_paid": row[8],
                "balance": row[9],
                "status": row[10],
            })
        return bills_data