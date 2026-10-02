import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from boarding_house import BoardingHouse
from owner import Owner
from tenant import Tenant
from single_room import SingleRoom
from shared_room import SharedRoom
from utility_reading import UtilityReading


class CleanupBackendTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'ledger.xlsx'
        self.house = BoardingHouse(Owner('Owner', ''), self.path)
        self.room = SingleRoom('1', 100)
        self.house.add_room(self.room)
        self.tenant = Tenant('Tenant One', '123')
        self.house.add_tenant(self.tenant, self.room)
        self.bill = self.house.generate_bills_for_room(UtilityReading(self.room, '2026-10', 3, 2, 10, 15), '2099-10-31')[0]

    def reload(self):
        return BoardingHouse(Owner('Owner', ''), self.path)

    def test_utility_components_survive_restart(self):
        self.assertEqual(getattr(self.bill, 'electricity_share', None), 30)
        bill = self.reload().bills[0]
        self.assertEqual(bill.electricity_share, 30)
        self.assertEqual(bill.water_share, 30)
        self.assertEqual(bill.get_total_amount(), 160)

    def test_payment_transactions_are_source_of_total_after_restart(self):
        self.house.record_payment(self.bill, 20, '2026-10-02')
        w = load_workbook(self.path)
        w['Billing'].cell(2, 8, 160)
        w['Billing'].cell(2, 10, 0)
        w['Billing'].cell(2, 11, 'PAID')
        w.save(self.path)
        w.close()
        bill = self.reload().bills[0]
        self.assertEqual(bill.amount_paid, 20)
        self.assertEqual(bill.balance, 140)
        self.assertEqual(bill.get_status(), 'PARTIAL')

    def test_edit_transfer_and_remove_preserve_history(self):
        self.assertTrue(hasattr(self.house, 'edit_tenant'))
        self.house.edit_tenant(self.tenant, 'Renamed Tenant', '456')
        target = SingleRoom('2', 200)
        self.house.add_room(target)
        self.house.transfer_tenant(self.tenant, target)
        house = self.reload()
        tenant = house.find_tenant('Renamed Tenant')
        self.assertEqual(tenant.assigned_room.room_number, '2')
        self.assertEqual(house.bills[0].room_number, '1')
        house.remove_tenant(tenant)
        house = self.reload()
        self.assertEqual(house.tenants, [])
        self.assertEqual(len(house.all_tenants), 1)
        self.assertEqual(house.bills[0].balance, 160)
        self.assertFalse(house.all_tenants[0].active)

    def test_room_edit_does_not_change_existing_bill(self):
        self.assertTrue(hasattr(self.house, 'edit_room'))
        self.house.edit_room(self.room, 500, 1)
        house = self.reload()
        self.assertEqual(house.rooms[0].base_rent, 500)
        self.assertEqual(house.bills[0].rent_share, 100)

    def test_external_change_is_preserved_and_failed_payment_rolls_back(self):
        workbook = load_workbook(self.path)
        workbook['Rooms'].cell(1, 5, 'Owner Notes')
        workbook['Rooms'].cell(2, 5, 'External edit')
        workbook.save(self.path)
        workbook.close()
        before = self.path.read_bytes()
        with self.assertRaisesRegex(OSError, 'changed outside'):
            self.house.record_payment(self.bill, 10, '2026-10-02')
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(self.bill.amount_paid, 0)
        self.assertEqual(self.house.payments, [])

    def test_migration_preserves_extra_columns_and_is_idempotent(self):
        workbook = load_workbook(self.path)
        workbook['Billing'].delete_cols(12, 3)
        workbook['Tenants'].delete_cols(4)
        del workbook['Payments']
        workbook['Billing'].cell(1, 12, 'Owner Notes')
        workbook['Billing'].cell(2, 12, 'Keep this')
        workbook['Tenants'].cell(1, 4, 'Custom')
        workbook['Tenants'].cell(2, 4, 'Keep tenant note')
        workbook.save(self.path)
        workbook.close()
        original = self.path.read_bytes()
        house = self.reload()
        self.assertEqual(house.ledger.backup_path.read_bytes(), original)
        workbook = load_workbook(self.path)
        self.assertEqual(workbook['Billing'].cell(2, 12).value, 'Keep this')
        self.assertEqual(workbook['Tenants'].cell(2, 4).value, 'Keep tenant note')
        workbook.close()
        migrated = self.path.read_bytes()
        again = self.reload()
        self.assertEqual(self.path.read_bytes(), migrated)
        self.assertIsNone(again.ledger.backup_path)
        self.assertFalse(again.bills[0].has_utility_breakdown)

    def test_shared_rounding_allocates_every_cent_once(self):
        room = SharedRoom('S', 100, 3)
        self.house.add_room(room)
        for i in range(3):
            self.house.add_tenant(Tenant(f'Shared {i}', '123'), room)
        bills = self.house.generate_bills_for_room(UtilityReading(room, '2026-10', 1, 1, 1, 1), '2099-10-31')
        self.assertEqual(round(sum(b.get_total_amount() for b in bills), 2), 102)
        self.assertEqual(round(sum(b.rent_share for b in bills), 2), 100)
        self.assertEqual(round(sum(b.electricity_share for b in bills), 2), 1)
        self.assertEqual(round(sum(b.water_share for b in bills), 2), 1)


if __name__ == '__main__':
    unittest.main()
