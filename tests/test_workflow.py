import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from boarding_house import BoardingHouse
from owner import Owner
from tenant import Tenant
from single_room import SingleRoom
from shared_room import SharedRoom
from utility_reading import UtilityReading


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'ledger.xlsx'
        self.house = BoardingHouse(Owner('Owner', '0917'), self.path)
        self.room = SingleRoom('101', 3500)
        self.house.add_room(self.room)
        self.tenant = Tenant('Test Tenant', '09171234567')
        self.house.add_tenant(self.tenant, self.room)
        self.bill = self.house.generate_bills_for_room(
            UtilityReading(self.room, '2026-10', 10, 2, 12, 30), '2099-10-31')[0]

    def reload(self):
        return BoardingHouse(Owner('Owner', '0917'), self.path)

    def test_partial_payment_is_not_paid_and_survives_restart(self):
        self.house.record_payment(self.bill, 1000, '2026-10-02')
        self.assertFalse(self.bill.paid)
        house = self.reload()
        self.assertEqual(house.bills[0].amount_paid, 1000)
        self.assertEqual(house.bills[0].balance, 2680)
        self.assertEqual(house.bills[0].get_status(), 'PARTIAL')
        house.record_payment(house.bills[0], 2680, '2026-10-03')
        house = self.reload()
        self.assertTrue(house.bills[0].paid)
        self.assertEqual([p.amount_paid for p in house.payments], [1000, 2680])
        self.assertEqual([p.remaining_balance for p in house.payments], [2680, 0])
        w = load_workbook(self.path)
        self.assertEqual([w['Billing'].cell(2, c).value for c in (8,10,11)], [3680,0,'PAID'])
        w.close()

    def test_reject_invalid_payments_without_changes(self):
        for amount in (-1, 0, 99999, float('nan'), float('inf'), 0.001):
            with self.subTest(amount=amount), self.assertRaises(ValueError):
                self.house.record_payment(self.bill, amount, '2026-10-02')
        with self.assertRaises(ValueError):
            self.house.record_payment(self.bill, 100, 'not-a-date')
        self.assertEqual(len(self.house.payments), 0)

    def test_duplicate_room_tenant_and_bill(self):
        with self.assertRaises(ValueError):
            self.house.add_room(SingleRoom('101', 1))
        room = SingleRoom('102', 1)
        self.house.add_room(room)
        with self.assertRaises(ValueError):
            self.house.add_tenant(Tenant('test tenant', '09170000123'), room)
        with self.assertRaises(ValueError):
            self.house.generate_bills_for_room(UtilityReading(self.room, '2026-10', 0, 0, 1, 1), '2099-10-31')

    def test_capacity_and_invalid_room_values(self):
        room = SharedRoom('S', 6000, 2)
        self.house.add_room(room)
        for name in ('One', 'Two'):
            self.house.add_tenant(Tenant(name, '09170000123'), room)
        with self.assertRaises(ValueError):
            self.house.add_tenant(Tenant('Three', '09170000123'), room)
        with self.assertRaises(ValueError):
            self.house.add_room(SharedRoom('bad', 100, 0))
        with self.assertRaises(ValueError):
            self.house.add_room(SingleRoom('negative', -1))

    def test_shared_rooms_hold_two_or_three_tenants(self):
        for capacity in (1, 4, 6):
            with self.subTest(capacity=capacity), self.assertRaises(ValueError):
                self.house.add_room(SharedRoom(f'S{capacity}', 6000, capacity))
        room = SharedRoom('S3', 6000, 3)
        self.house.add_room(room)
        with self.assertRaises(ValueError):
            self.house.edit_room(room, 6000, 4)
        self.house.edit_room(room, 6000, 2)
        self.assertEqual(self.reload().find_room('S3').get_max_occupancy(), 2)

    def test_contact_number_must_be_eleven_digits(self):
        room = SharedRoom('S', 6000, 3)
        self.house.add_room(room)
        for contact in ('0917123456', '091712345678', '0917-123-4567', '0917123456a', '0981231879023'):
            with self.subTest(contact=contact), self.assertRaises(ValueError):
                self.house.add_tenant(Tenant(f'Tenant {contact}', contact), room)
            with self.subTest(edit=contact), self.assertRaises(ValueError):
                self.house.edit_tenant(self.tenant, self.tenant.name, contact)
        self.house.add_tenant(Tenant('Valid Tenant', ' 09171234567 '), room)
        self.assertEqual(self.reload().find_tenant('Valid Tenant').contact_number, '09171234567')
        self.assertEqual(self.tenant.contact_number, '09171234567')

    def test_dates_and_months_require_exact_format(self):
        for value in ('2026/10/02', '20261002', '2026-10-2', '10-02-2026', '2026-13-01', '2026-02-30'):
            with self.subTest(date=value), self.assertRaises(ValueError):
                self.house.record_payment(self.bill, 100, value)
        for month in ('2026-1', '202611', '2026/11', '2026-11-01', '11-2026'):
            with self.subTest(month=month), self.assertRaises(ValueError):
                self.house.generate_bills_for_room(UtilityReading(self.room, month, 0, 0, 1, 1), '2099-11-30')
        with self.assertRaises(ValueError):
            self.house.generate_bills_for_room(UtilityReading(self.room, '2026-11', 0, 0, 1, 1), '2099/11/30')
        self.assertEqual(len(self.house.payments), 0)
        self.assertEqual(len(self.house.bills), 1)

    def test_failed_save_does_not_apply_payment(self):
        with patch.object(self.house.ledger, '_save', side_effect=PermissionError('locked')):
            with self.assertRaises(PermissionError):
                self.house.record_payment(self.bill, 100, '2026-10-02')
        self.assertFalse(self.bill.paid)
        self.assertEqual(len(self.house.payments), 0)
        self.assertEqual(self.reload().bills[0].amount_paid, 0)
        self.house.record_payment(self.bill, 100, '2026-10-02')
        self.assertEqual(self.reload().bills[0].amount_paid, 100)

    def test_overdue_partial_and_paid(self):
        self.house.record_payment(self.bill, 100, '2026-10-02')
        self.assertEqual(self.bill.get_status('2099-11-01'), 'OVERDUE')
        self.house.record_payment(self.bill, self.bill.balance, '2026-10-03')
        self.assertEqual(self.bill.get_status('2099-11-01'), 'PAID')

    def test_names_that_start_with_equals_remain_text_after_restart(self):
        room = SingleRoom('=Suite A', 100)
        self.house.add_room(room)
        tenant = Tenant('=Tenant Name', '09170000000')
        self.house.add_tenant(tenant, room)
        self.house.generate_bills_for_room(UtilityReading(room, '2026-11', 0, 0, 0, 0), '2099-11-30')
        workbook = load_workbook(self.path)
        self.assertEqual(workbook['Tenants'].cell(3, 1).data_type, 's')
        self.assertEqual(workbook['Rooms'].cell(3, 1).data_type, 's')
        workbook.close()
        self.assertIsNotNone(self.reload().find_tenant('=Tenant Name'))

    def test_reminders_after_restart_use_remaining_balance_and_exclude_paid(self):
        self.house.record_payment(self.bill, 1000, '2026-10-02')
        house = self.reload()
        reminders = house.send_all_reminders(send=False)
        self.assertEqual(len(reminders), 1)
        self.assertIn('2680.00', reminders[0].message)
        house.record_payment(house.bills[0], 2680, '2026-10-03')
        self.assertEqual(house.send_all_reminders(send=False), [])

    def test_failed_room_and_batch_bill_save_leave_no_phantom_records(self):
        room = SharedRoom('202', 6000, 2)
        with patch.object(self.house.ledger, '_save', side_effect=PermissionError('locked')):
            with self.assertRaises(PermissionError):
                self.house.add_room(room)
        self.assertIsNone(self.house.find_room('202'))
        self.house.add_room(room)
        self.house.add_tenant(Tenant('One', '09170000123'), room)
        self.house.add_tenant(Tenant('Two', '09170000456'), room)
        with patch.object(self.house.ledger, '_save', side_effect=PermissionError('locked')):
            with self.assertRaises(PermissionError):
                self.house.generate_bills_for_room(UtilityReading(room, '2026-10', 0, 0, 0, 0), '2099-10-31')
        self.assertEqual(len(self.house.bills), 1)
        self.assertEqual(len(self.reload().bills), 1)

    def test_existing_workbook_migration_preserves_source_and_money(self):
        source = Path(__file__).resolve().parents[1] / 'billing_ledger.xlsx'
        destination = Path(self.temp.name) / 'existing.xlsx'
        shutil.copy2(source, destination)
        # Exercise migration even after the real workbook has already been upgraded.
        fixture = load_workbook(destination)
        if 'Payments' in fixture.sheetnames:
            del fixture['Payments']
        if fixture['Billing'].max_column > 11:
            fixture['Billing'].delete_cols(12, fixture['Billing'].max_column-11)
        if fixture['Tenants'].max_column > 3:
            fixture['Tenants'].delete_cols(4, fixture['Tenants'].max_column-3)
        fixture.save(destination)
        fixture.close()
        original = destination.read_bytes()
        old = load_workbook(destination)
        expected = sum(float(r[7] or 0) for r in list(old['Billing'].values)[1:])
        old.close()
        house = BoardingHouse(Owner('Owner', ''), destination)
        self.assertFalse(house.bills[0].paid)
        self.assertEqual(sum(b.amount_paid for b in house.bills), expected)
        backups = list((destination.parent / 'backups').glob('*.xlsx'))
        self.assertTrue(any(p.read_bytes() == original for p in backups))
        again = BoardingHouse(Owner('Owner', ''), destination)
        self.assertEqual(sum(p.amount_paid for p in again.payments), expected)


if __name__ == '__main__':
    unittest.main()
